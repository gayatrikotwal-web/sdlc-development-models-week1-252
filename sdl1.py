
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Flowable
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.enums import TA_CENTER
import os

OUTPUT_DIR = os.path.abspath(".")
DOCX_FILE = os.path.join(
    OUTPUT_DIR, "STLC_Assignment_Submission.docx"
)
PDF_FILE = os.path.join(
    OUTPUT_DIR, "STLC_Flow_Diagram.pdf"
)

# --------------------------------------------------
# STLC PHASE CONTENT
# --------------------------------------------------

phases = [
    {
        "title": "1. Requirement Analysis",
        "description": (
            "The testing team reviews business requirements, "
            "user stories, functional and non-functional requirements, "
            "and acceptance criteria. The team identifies testable "
            "requirements, risks, dependencies, and unclear areas."
        ),
        "entry": [
            "Business requirements or user stories are available.",
            "Project objectives and scope are sufficiently understood.",
            "Stakeholders are available to clarify requirements."
        ],
        "exit": [
            "Requirements are reviewed for testability.",
            "Ambiguities and questions are documented.",
            "Initial test conditions are identified.",
            "Requirement-to-test traceability is prepared where possible."
        ],
        "deliverables": (
            "Reviewed requirements, test conditions, clarification log, "
            "and initial Requirements Traceability Matrix (RTM)."
        )
    },
    {
        "title": "2. Test Planning",
        "description": (
            "The test lead defines testing scope, objectives, strategy, "
            "resources, schedule, tools, risks, responsibilities, "
            "and reporting methods. Estimates and deliverables are agreed."
        ),
        "entry": [
            "Requirement analysis is sufficiently complete.",
            "Project scope and priorities are understood.",
            "Resource and schedule constraints can be identified."
        ],
        "exit": [
            "The test plan is reviewed and approved.",
            "Scope, schedule, responsibilities, and estimates are documented.",
            "Tools, risks, resources, and deliverables are identified.",
            "Entry and exit criteria are agreed upon."
        ],
        "deliverables": (
            "Test plan, test strategy, estimates, schedule, "
            "resource plan, and risk assessment."
        )
    },
    {
        "title": "3. Test Case Development",
        "description": (
            "Testers prepare test scenarios, detailed test cases, test "
            "data, and expected results. Test cases are reviewed and "
            "mapped to requirements. Automation scripts may also be prepared."
        ),
        "entry": [
            "Requirements are stable enough for test design.",
            "Test conditions and acceptance criteria are available.",
            "The testing approach is understood."
        ],
        "exit": [
            "Test cases and expected results are documented and reviewed.",
            "Required test data is prepared or its setup is defined.",
            "Requirement-to-test traceability is established.",
            "Automation scripts are reviewed if automation is in scope."
        ],
        "deliverables": (
            "Test scenarios, test cases, expected results, test data, "
            "and automation scripts where applicable."
        )
    },
    {
        "title": "4. Test Environment Setup",
        "description": (
            "The team prepares the hardware, software, network, accounts, "
            "configurations, services, and test data required for testing. "
            "The environment is verified before formal execution."
        ),
        "entry": [
            "Environment requirements and configuration needs are known.",
            "Required infrastructure and dependencies can be provisioned.",
            "A testable application build can be provided."
        ],
        "exit": [
            "The test environment is configured.",
            "The required application build is deployed.",
            "Access, integrations, and test data are verified.",
            "Smoke tests confirm basic environment readiness.",
            "Environment blockers are recorded."
        ],
        "deliverables": (
            "Configured test environment, build details, smoke-test "
            "results, and environment readiness confirmation."
        )
    },
    {
        "title": "5. Test Execution",
        "description": (
            "Testers execute test cases and compare actual results with "
            "expected results. They record outcomes, report defects, "
            "retest fixes, and perform regression testing when required."
        ),
        "entry": [
            "Approved test cases are available.",
            "The test environment is ready.",
            "Test data is prepared.",
            "A testable application build is deployed.",
            "Required initial checks, such as smoke tests, pass."
        ],
        "exit": [
            "Planned tests are executed or exceptions are documented.",
            "Pass, fail, and blocked results are recorded.",
            "Defects are logged and prioritized.",
            "Fixed defects are retested.",
            "Regression testing is performed where required.",
            "Remaining risks and unresolved defects are documented."
        ],
        "deliverables": (
            "Execution report, defect reports, test evidence, "
            "retest results, and regression-test results."
        )
    },
    {
        "title": "6. Test Cycle Closure",
        "description": (
            "The team evaluates testing against objectives and exit "
            "criteria. It summarizes coverage and defect status, "
            "documents residual risks, captures lessons learned, "
            "and prepares a test summary report."
        ),
        "entry": [
            "Test execution results and defect records are available.",
            "Test coverage information is updated.",
            "The cycle is complete or formally stopped.",
            "Results are ready for stakeholder review."
        ],
        "exit": [
            "The test summary report is prepared and shared.",
            "Coverage and pass/fail results are documented.",
            "Open defects, limitations, and residual risks are recorded.",
            "Lessons learned are captured.",
            "Closure or any release-risk decision is documented."
        ],
        "deliverables": (
            "Test summary report, coverage and defect metrics, "
            "residual-risk report, and lessons learned."
        )
    }
]


# --------------------------------------------------
# CREATE THE WORD ASSIGNMENT
# --------------------------------------------------

def setup_word_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    for name, size in [
        ("Title", 25),
        ("Heading 1", 16),
        ("Heading 2", 12)
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def create_word_document():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    setup_word_styles(doc)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SDLC & DEVELOPMENT MODELS")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(89, 89, 89)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Assignment: Map the STLC")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Software Testing Life Cycle")
    run.italic = True
    run.font.size = Pt(13)

    doc.add_paragraph("")
    for label in [
        "Name: ______________________________",
        "Class / Team: ________________________",
        "Date: _______________________________"
    ]:
        p = doc.add_paragraph(label)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "The Software Testing Life Cycle (STLC) is a structured "
        "sequence of testing activities used to assess software quality. "
        "Each phase has objectives, deliverables, entry criteria, and "
        "exit criteria. Entry criteria define the conditions expected "
        "before a phase begins; exit criteria define the conditions "
        "expected before a phase is completed. Criteria should be "
        "adapted to project size, risk, and development approach."
    )

    # Flow summary
    doc.add_heading("2. STLC Flow", level=1)
    doc.add_paragraph(
        "Requirement Analysis → Test Planning → Test Case Development "
        "→ Test Environment Setup → Test Execution → Test Cycle Closure"
    )
    doc.add_paragraph(
        "STLC is iterative. When a defect is found, it is reported to "
        "development. After a fix, testers retest and may perform "
        "regression testing. Changed requirements or unmet exit criteria "
        "may require revisiting an earlier phase."
    )

    # Phase details
    doc.add_heading(
        "3. STLC Phases with Entry and Exit Criteria", level=1
    )

    for phase in phases:
        doc.add_heading(phase["title"], level=2)
        doc.add_paragraph(phase["description"])

        p = doc.add_paragraph()
        run = p.add_run("Entry Criteria")
        run.bold = True
        add_bullet_list(doc, phase["entry"])

        p = doc.add_paragraph()
        run = p.add_run("Exit Criteria")
        run.bold = True
        add_bullet_list(doc, phase["exit"])

        p = doc.add_paragraph()
        run = p.add_run("Deliverables: ")
        run.bold = True
        p.add_run(phase["deliverables"])

    # Deliverables table
    doc.add_heading("4. Deliverables Summary", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "STLC Phase"
    table.rows[0].cells[1].text = "Typical Deliverables"

    for phase in phases:
        cells = table.add_row().cells
        cells[0].text = phase["title"]
        cells[1].text = phase["deliverables"]

    # Important notes
    doc.add_heading("5. Important Notes", level=1)
    add_bullet_list(doc, [
        "Entry and exit criteria should be agreed upon before testing "
        "begins and tailored to the project.",
        "A defect does not automatically stop all testing. Severity, "
        "business risk, and agreed thresholds guide decisions.",
        "In Agile projects, STLC activities can overlap and repeat "
        "within each sprint.",
        "Testing supports the wider SDLC and complements reviews "
        "and quality practices throughout development."
    ])

    # Conclusion
    doc.add_heading("6. Conclusion", level=1)
    doc.add_paragraph(
        "STLC provides a repeatable approach for planning, designing, "
        "executing, and closing software tests. Clear entry and exit "
        "criteria improve readiness, traceability, defect management, "
        "and communication. A test summary report helps stakeholders "
        "understand software quality and remaining risks."
    )

    doc.save(DOCX_FILE)
    print(f"Word assignment created: {DOCX_FILE}")


# --------------------------------------------------
# CREATE THE STLC FLOW DIAGRAM PDF
# --------------------------------------------------

class STLCFlowDiagram(Flowable):
    def __init__(self, width):
        super().__init__()
        self.width = width
        self.height = 320

    def draw(self):
        canvas = self.canv
        width = self.width

        boxes = [
            ("1. REQUIREMENT\nANALYSIS",
             "Understand and assess\nrequirements"),
            ("2. TEST\nPLANNING",
             "Define scope, strategy,\nresources and schedule"),
            ("3. TEST CASE\nDEVELOPMENT",
             "Design cases, test data\nand expected results"),
            ("4. TEST ENVIRONMENT\nSETUP",
             "Configure and verify\nthe test environment"),
            ("5. TEST\nEXECUTION",
             "Run tests, log defects,\nretest and regress"),
            ("6. TEST CYCLE\nCLOSURE",
             "Report results, risks\nand lessons learned")
        ]

        gap = 12
        box_width = (width - 5 * gap) / 6
        box_height = 105
        box_y = 180

        for index, (heading, description) in enumerate(boxes):
            x = index * (box_width + gap)

            canvas.setFillColor(
                colors.HexColor("#EAF2F8")
                if index % 2 == 0
                else colors.HexColor("#F4F7FA")
            )
            canvas.setStrokeColor(colors.HexColor("#4472C4"))
            canvas.roundRect(
                x, box_y, box_width, box_height,
                7, stroke=1, fill=1
            )

            canvas.setFillColor(colors.HexColor("#1F4E79"))
            canvas.setFont("Helvetica-Bold", 7.8)

            y = box_y + box_height - 23
            for line in heading.split("\n"):
                canvas.drawCentredString(
                    x + box_width / 2, y, line
                )
                y -= 10

            canvas.setStrokeColor(colors.HexColor("#A6B6C7"))
            canvas.line(
                x + 7, y - 2,
                x + box_width - 7, y - 2
            )

            canvas.setFillColor(colors.HexColor("#333333"))
            canvas.setFont("Helvetica", 7.1)
            for line_index, line in enumerate(
                description.split("\n")
            ):
                canvas.drawCentredString(
                    x + box_width / 2,
                    y - 15 - line_index * 10,
                    line
                )

            # Arrow between phases
            if index < len(boxes) - 1:
                start_x = x + box_width + 1
                middle_y = box_y + box_height / 2

                canvas.setStrokeColor(colors.HexColor("#527DAA"))
                canvas.setFillColor(colors.HexColor("#527DAA"))
                canvas.setLineWidth(1.4)
                canvas.line(
                    start_x, middle_y,
                    start_x + gap - 4, middle_y
                )

                arrow = canvas.beginPath()
                arrow.moveTo(
                    start_x + gap - 9, middle_y + 4
                )
                arrow.lineTo(
                    start_x + gap - 3, middle_y
                )
                arrow.lineTo(
                    start_x + gap - 9, middle_y - 4
                )
                arrow.close()
                canvas.drawPath(arrow, stroke=0, fill=1)

        # Feedback loop
        loop_y = 140
        x_start = width * 0.77
        x_end = width * 0.56

        canvas.setStrokeColor(colors.HexColor("#C0504D"))
        canvas.setLineWidth(1.4)
        canvas.line(x_start, box_y - 2, x_start, loop_y)
        canvas.line(x_start, loop_y, x_end, loop_y)
        canvas.line(x_end, loop_y, x_end, box_y - 2)

        canvas.setFillColor(colors.HexColor("#9C2F2F"))
        canvas.setFont("Helvetica-Bold", 8)
        canvas.drawCentredString(
            (x_start + x_end) / 2,
            loop_y - 13,
            "Defect fix → retest → regression testing"
        )

        canvas.setFillColor(colors.HexColor("#333333"))
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(
            width / 2, 82,
            "Entry criteria confirm readiness; exit criteria "
            "confirm required work is complete."
        )
        canvas.drawCentredString(
            width / 2, 64,
            "If criteria are not met, revisit the relevant phase."
        )

        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawCentredString(
            width / 2, 34,
            "STLC is iterative; activities may overlap in Agile projects."
        )


def create_flow_diagram():
    page_size = landscape(A4)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "STLCTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1F4E79")
    )
    subtitle_style = ParagraphStyle(
        "STLCSubtitle",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#555555")
    )

    def footer(canvas, pdf_doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D9E2F3"))
        canvas.line(
            30, 25, page_size[0] - 30, 25
        )
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(
            30, 12, "Assignment: Map the STLC"
        )
        canvas.drawRightString(
            page_size[0] - 30, 12,
            f"Page {pdf_doc.page}"
        )
        canvas.restoreState()

    story = [
        Paragraph(
            "SOFTWARE TESTING LIFE CYCLE (STLC)",
            title_style
        ),
        Paragraph(
            "SDLC & Development Models — Assignment Diagram",
            subtitle_style
        ),
        Spacer(1, 22),
        STLCFlowDiagram(page_size[0] - 70)
    ]

    pdf = SimpleDocTemplate(
        PDF_FILE,
        pagesize=page_size,
        rightMargin=35,
        leftMargin=35,
        topMargin=32,
        bottomMargin=35
    )
    pdf.build(
        story,
        onFirstPage=footer,
        onLaterPages=footer
    )
    print(f"Flow diagram created: {PDF_FILE}")


# --------------------------------------------------
# MAIN PROGRAM
# --------------------------------------------------

if __name__ == "__main__":
    create_word_document()
    create_flow_diagram()

    print("\nAssignment generation completed.")
    print("Files saved in:", OUTPUT_DIR)
